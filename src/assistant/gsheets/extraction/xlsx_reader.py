import io
import re
import warnings
from dataclasses import dataclass
from typing import Optional

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from .models import Cell

warnings.filterwarnings(
    "ignore",
    category=UserWarning,
    module="openpyxl.worksheet._reader",
)


XLSX_MIME = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

_THREADED_PREFIX = re.compile(
    r"^\[Threaded comment\].*?Comment:\s*",
    re.DOTALL,
)

@dataclass
class SheetMatrix:
    name: str
    rows: list[list[Cell]]


def download_xlsx(
    spreadsheet_id: str,
    credentials_path: str,
    max_retries: int = 5,
) -> bytes:

    import time

    from google.oauth2.service_account import Credentials
    from googleapiclient.discovery import build
    from googleapiclient.errors import HttpError

    creds = Credentials.from_service_account_file(
        credentials_path,
        scopes=["https://www.googleapis.com/auth/drive.readonly"],
    )
    drive = build("drive", "v3", credentials=creds, cache_discovery=False)

    last_error = None
    for attempt in range(1, max_retries + 1):
        try:
            return drive.files().export_media(
                fileId=spreadsheet_id,
                mimeType=XLSX_MIME,
            ).execute()
        except HttpError as e:
            # Повтор для ошибок google
            if e.resp.status < 500 or attempt == max_retries:
                raise
            last_error = e
            delay = 2 ** (attempt - 1)
            print(
                f"  Google вернул {e.resp.status}, "
                f"попытка {attempt}/{max_retries}, "
                f"повтор через {delay} сек..."
            )
            time.sleep(delay)

    raise last_error

def _extract_value(cell) -> str:
    v = cell.value
    if v is None:
        return ""

    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return str(v)

    if isinstance(v, int):
        return str(v)

    if hasattr(v, "strftime"):
        try:
            return v.strftime("%Y-%m-%d")
        except Exception:
            return str(v)

    return str(v).strip()


def _extract_color(cell) -> Optional[str]:
    fill = cell.fill
    if not fill or fill.patternType is None:
        return None

    color = getattr(fill, "start_color", None) or getattr(fill, "fgColor", None)
    if color is None:
        return None

    rgb = color.rgb
    if not isinstance(rgb, str):
        return None

    if len(rgb) == 8:
        r, g, b = rgb[2:4], rgb[4:6], rgb[6:8]
        if rgb in ("FFFFFFFF", "00FFFFFF"):
            return None
        return f"#{r}{g}{b}".upper()

    if len(rgb) == 6:
        if rgb.upper() == "FFFFFF":
            return None
        return f"#{rgb}".upper()

    return None


def _extract_comment(cell) -> str:
    if not cell.comment:
        return ""
    text = cell.comment.text or ""
    text = _THREADED_PREFIX.sub("", text).strip()
    return text



def _find_real_bounds(ws: Worksheet, max_rows: int = 2000, max_cols: int = 100):
    last_row = 0
    last_col = 0
    limit_rows = min(max_rows, ws.max_row)
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=limit_rows), start=1):
        for c_idx, cell in enumerate(row, start=1):
            if c_idx > max_cols:
                break
            if cell.value is not None and str(cell.value).strip():
                if r_idx > last_row:
                    last_row = r_idx
                if c_idx > last_col:
                    last_col = c_idx
    return last_row, last_col


def parse_workbook(data: bytes) -> list[SheetMatrix]:
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)

    matrices: list[SheetMatrix] = []
    for name in wb.sheetnames:
        ws = wb[name]
        last_row, last_col = _find_real_bounds(ws)

        rows: list[list[Cell]] = []
        for r in range(1, last_row + 1):
            row_cells: list[Cell] = []
            for c in range(1, last_col + 1):
                cell = ws.cell(row=r, column=c)
                row_cells.append(Cell(
                    value=_extract_value(cell),
                    color=_extract_color(cell),
                    comment=_extract_comment(cell),
                ))
            rows.append(row_cells)

        matrices.append(SheetMatrix(name=name, rows=rows))

    return matrices



def load_workbook_matrices(
    spreadsheet_id: str,
    credentials_path: str,
) -> list[SheetMatrix]:
    data = download_xlsx(spreadsheet_id, credentials_path)
    return parse_workbook(data)